"""Ground truth and scoring live here; the inference pipeline never imports them."""

import hashlib
import math
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from backend.codex import Codex, normalize_clip


def parse_timestamp(value: str | None) -> float | None:
    if value is None or not str(value).strip():
        return None
    text = str(value).lower().strip().replace(",", ".")
    match = re.fullmatch(
        r"(?:(\d+)\s*(?:минута|минуты|минут|мин|м)\s*)?"
        r"(?:(\d+(?:\.\d+)?)\s*(?:секунда|секунды|секунд|сек|с)\s*)?",
        text,
    )
    if not match or not any(match.groups()):
        raise ValueError(f"Invalid label timestamp: {value!r}")
    minutes, seconds = match.groups()
    if minutes is not None and seconds is not None and float(seconds) >= 60:
        raise ValueError(f"Invalid seconds component: {value!r}")
    result = float(minutes or 0) * 60 + float(seconds or 0)
    if not math.isfinite(result):
        raise ValueError(f"Non-finite label timestamp: {value!r}")
    return result


@dataclass(frozen=True)
class LabelEvent:
    timestamp_sec: float
    rule_ids: tuple[str, ...]


@dataclass
class VideoGroup:
    path: Path
    checksum: str
    aliases: list[str] = field(default_factory=list)
    labels: list[LabelEvent] = field(default_factory=list)
    layouts: set[str] = field(default_factory=set)


def read_labels(path: Path) -> dict[str, list[float]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = iter(workbook.active.iter_rows(values_only=True))
        header = [str(v).strip().lower() for v in next(rows)]
        required = ["наименование файла", "таймстемп", "нарушение / без нарушения"]
        if not all(name in header for name in required):
            raise ValueError("Label workbook is missing required columns")
        indices = [header.index(name) for name in required]
        result: dict[str, list[float]] = {}
        for row_number, row in enumerate(rows, start=2):
            if not any(v is not None for v in row):
                continue
            name, raw_time, raw_verdict = [row[i] for i in indices]
            clip = normalize_clip(name)
            timestamp = parse_timestamp(raw_time)
            verdict = str(raw_verdict).strip().lower()
            if verdict not in {"нарушение", "без нарушения"}:
                raise ValueError(f"Unknown verdict at label row {row_number}: {raw_verdict!r}")
            if (timestamp is None) != (verdict == "без нарушения"):
                raise ValueError(f"Timestamp/verdict conflict at label row {row_number}")
            events = result.setdefault(clip, [])
            if timestamp is not None and timestamp not in events:
                events.append(timestamp)
        if not result:
            raise ValueError("Label workbook contains no clips")
        return result
    finally:
        workbook.close()


def group_videos(videos: Path, labels: dict[str, list[float]], codex: Codex) -> list[VideoGroup]:
    paths = sorted(p for p in videos.iterdir() if p.is_file() and p.suffix.lower() == ".mov")
    if not paths:
        raise ValueError(f"No .mov videos found in {videos}")
    metadata = {clip.id: clip for clip in codex.clips}
    groups: dict[str, VideoGroup] = {}
    found: set[str] = set()
    for path in paths:
        clip = normalize_clip(path.name)
        if clip in found:
            raise ValueError(f"Several files map to the same clip: {clip}")
        if clip not in labels:
            raise ValueError(f"Video has no label: {path.name}")
        found.add(clip)
        with path.open("rb") as stream:
            checksum = hashlib.file_digest(stream, "sha256").hexdigest()
        group = groups.setdefault(checksum, VideoGroup(path=path, checksum=checksum))
        group.aliases.append(clip)
        info = metadata.get(clip)
        if info is not None:
            group.layouts.add(info.camera_layout)
        for timestamp in labels[clip]:
            rule_ids = tuple(info.rule_ids) if info else ()
            existing = next((e for e in group.labels if e.timestamp_sec == timestamp), None)
            if existing:
                group.labels.remove(existing)
                rule_ids = tuple(sorted(set(existing.rule_ids) | set(rule_ids)))
            group.labels.append(LabelEvent(timestamp, rule_ids))
    if missing := set(labels) - found:
        raise ValueError(f"Labelled clips are missing from videos: {', '.join(sorted(missing))}")
    for group in groups.values():
        group.labels.sort(key=lambda event: event.timestamp_sec)
    return list(groups.values())


def match_events(
    labels: list[float], predictions: list[float], tolerance: float
) -> list[tuple[int, int]]:
    """Maximum one-to-one temporal matching; a prediction cannot hit two labels."""
    if not math.isfinite(tolerance) or tolerance < 0:
        raise ValueError("Tolerance must be non-negative and finite")
    if any(not math.isfinite(t) or t < 0 for t in labels + predictions):
        raise ValueError("Event times must be non-negative and finite")
    left, right = (
        sorted(enumerate(labels), key=lambda x: x[1]),
        sorted(enumerate(predictions), key=lambda x: x[1]),
    )
    i = j = 0
    pairs = []
    while i < len(left) and j < len(right):
        if abs(left[i][1] - right[j][1]) <= tolerance:
            pairs.append((left[i][0], right[j][0]))
            i += 1
            j += 1
        elif left[i][1] < right[j][1]:
            i += 1
        else:
            j += 1
    return pairs


def score(rows: list[dict], codex: Codex, scope: str, *, weight_aliases: bool = False) -> dict:
    def weight(row: dict) -> int:
        return len(row["aliases"]) if weight_aliases else 1

    count = sum(weight(row) for row in rows)
    positives = sum(weight(row) for row in rows if row["labels"])
    negatives = count - positives
    tp = tn = fp = failed = 0
    hits = {1: 0, 2: 0, 5: 0}
    event_count = prediction_count = 0
    per_rule = {rule.id: {"hits": 0, "misses": 0, "false_alarms": 0} for rule in codex.rules}
    for row in rows:
        w = weight(row)
        labels = [event["timestamp_sec"] for event in row["labels"]]
        predictions = row["result"]["violations"] if row["result"] else []
        times = [v["timecode_sec"] for v in predictions]
        event_count += len(labels) * w
        prediction_count += len(times) * w
        if row["status"] == "failed":
            failed += w
        elif labels:
            tp += bool(times) * w
        else:
            fp += bool(times) * w
            tn += (not times) * w
        for tolerance in hits:
            hits[tolerance] += len(match_events(labels, times, tolerance)) * w
        for rule_id, counts in per_rule.items():
            rule_labels = [e["timestamp_sec"] for e in row["labels"] if rule_id in e["rule_ids"]]
            rule_times = [v["timecode_sec"] for v in predictions if v["rule_id"] == rule_id]
            matched = len(match_events(rule_labels, rule_times, 2))
            counts["hits"] += matched * w
            counts["misses"] += (len(rule_labels) - matched) * w
            counts["false_alarms"] += (len(rule_times) - matched) * w
    slowest = max(rows, key=lambda row: row["end_to_end_sec"], default=None)
    return {
        "scope": scope,
        "input_clips": sum(len(r["aliases"]) for r in rows),
        "unique_videos": len(rows),
        "scored_videos": count,
        "violation_videos": positives,
        "correct_videos": negatives,
        "failed": failed,
        "accuracy": (tp + tn) / count if count else None,
        "recall": tp / positives if positives else None,
        "false_alarm_rate": fp / negatives if negatives else None,
        "labelled_events": event_count,
        "predicted_events": prediction_count,
        **{f"hits_{t}s": n for t, n in hits.items()},
        **{f"event_recall_{t}s": n / event_count if event_count else None for t, n in hits.items()},
        "false_events_2s": prediction_count - hits[2],
        "unmapped_label_events": sum(
            weight(row) for row in rows for event in row["labels"] if not event["rule_ids"]
        ),
        "mean_runtime_sec": sum(row["end_to_end_sec"] * weight(row) for row in rows) / count
        if count
        else None,
        "slowest_video": "/".join(slowest["aliases"]) if slowest else None,
        "slowest_sec": slowest["end_to_end_sec"] if slowest else None,
        "over_300s": sum(weight(row) for row in rows if row["end_to_end_sec"] > 300),
        "per_rule": per_rule,
    }
