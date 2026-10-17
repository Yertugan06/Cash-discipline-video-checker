"""Build the private ground-truth file used by every experiment.

Sources (kept separate in the output):
  organizer : data/open/open_labels.xlsx          -> verdict + integer-second time (the only ground truth)
  human     : data/Таблица_нарушений.xlsx          -> human event seconds, norm codes, confidence, review status
  layout    : data/Cash_Discipline_Cases_Rules.xlsx -> camera-layout tag per clip (provenance: team workbook)

Writes validation-private/gt.json (gitignored). Prints aggregate counts only.
Usage: python -I gt.py <research_root>   (e.g. validation; organizer data is read from <research_root>/../data)
"""
import collections, hashlib, json, os, re, sys

import openpyxl

ROOT = sys.argv[1]
OUT = os.path.join(ROOT, 'validation-private')
UNRELIABLE = {'X-1', 'X-11', 'X-13', 'X-44'}   # flagged by the team as uncertain (user, 2026-10-17)
REVIEWED_OK = {'X-24', 'X-32'}                 # manually re-reviewed and agreed (user, 2026-10-17)


def norm_id(s):
    s = str(s).strip().replace('Х', 'X').replace('х', 'X').replace('x', 'X')
    m = re.search(r'X\s*-?\s*(\d+)', s)
    return f'X-{int(m.group(1))}' if m else None


def seconds(s):
    """'1минута 8 секунда' -> [68]; '57' -> [57]; '57 / 68' -> [57, 68]."""
    s = str(s)
    if 'мин' in s:
        n = [int(x) for x in re.findall(r'\d+', s)]
        return [n[0] * 60 + (n[1] if len(n) > 1 else 0)]
    return [float(x.replace(',', '.')) for x in re.findall(r'\d+(?:[.,]\d+)?', s)]


def sheet_rows(path, name, skip):
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)[name]
    rows = [r for r in ws.iter_rows(values_only=True) if any(v not in (None, '') for v in r)]
    return rows[skip:]


# ---- organizer labels + files ---------------------------------------------------------
vdir = os.path.join(ROOT, os.pardir, 'data', 'open', 'videos')
files = {norm_id(f): os.path.join(vdir, f) for f in os.listdir(vdir) if f.endswith('.mov')}
md5 = {k: hashlib.md5(open(p, 'rb').read()).hexdigest() for k, p in files.items()}
org = {}
wb =openpyxl.load_workbook(os.path.join(ROOT, os.pardir, 'data', 'open', 'open_labels.xlsx'), read_only=True, data_only=True)
for r in list(wb.worksheets[0].iter_rows(values_only=True))[1:]:
    if not r or not r[0]:
        continue
    cid = norm_id(r[0])
    is_v = any('наруш' in str(x).lower() and 'без' not in str(x).lower() for x in r[1:] if x)
    t = [x for c in r[1:] if c and re.search(r'\d', str(c)) for x in seconds(c)]
    org[cid] = {'violation': bool(t) or is_v, 'times': t}

# ---- duplicate groups -----------------------------------------------------------------
by_md5 = collections.defaultdict(list)
for k in sorted(files, key=lambda s: int(s[2:])):
    by_md5[md5[k]].append(k)
canon = {k: ks[0] for ks in by_md5.values() for k in ks}

# ---- human annotations ----------------------------------------------------------------
hp = os.path.join(ROOT, os.pardir, 'data', 'Таблица_нарушений.xlsx')
human = collections.defaultdict(list)
pairs = []
for r in sheet_rows(hp, 'Нарушения', 3):
    cid = norm_id(r[1]) if r[1] else None
    if not cid:
        continue
    code = re.sub(r'^открыть\s*', '', str(r[7] or '')).strip() or None
    human[cid].append({'times': seconds(r[2]) if r[2] is not None else [], 'type': r[5], 'norm1': code,
                       'confidence': r[11], 'must_check': r[12], 'review': r[15]})
    for p in re.findall(r'[XХxх]\s*-\s*\d+', str(r[14] or '')):
        pairs.append((cid, norm_id(p)))
clean_rows = 0
for r in sheet_rows(hp, 'Без нарушений', 1):
    cid = norm_id(r[0]) if r[0] else None
    if not cid:
        continue
    clean_rows += 1
    for p in re.findall(r'[XХxх]\s*-\s*\d+', str(r[2] or '')):
        pairs.append((cid, norm_id(p)))

# ---- camera layout --------------------------------------------------------------------
layout = {}
for r in sheet_rows(os.path.join(ROOT, os.pardir, 'data', 'Cash_Discipline_Cases_Rules.xlsx'), 'Clips', 1):
    if r[0]:
        layout[norm_id(r[0])] = r[8]

# ---- groups = duplicates + violation/clean pairs (union-find over canonical ids) ------
parent = {}
def find(x):
    parent.setdefault(x, x)
    while parent[x] != x:
        parent[x] = parent[parent[x]]; x = parent[x]
    return x
for a, b in pairs:
    if a in canon and b in canon:
        parent[find(canon[a])] = find(canon[b])

clips = {}
for k in sorted(set(canon.values()), key=lambda s: int(s[2:])):
    dups = [d for d in canon if canon[d] == k]
    o_times = sorted({t for d in dups for t in org[d]['times']})
    h_rows = [h for d in dups for h in human.get(d, [])]
    reliable = bool(h_rows) and not (set(dups) & UNRELIABLE) and all(
        str(h['must_check']) != 'ОБЯЗАТЕЛЬНО' or set(dups) & REVIEWED_OK for h in h_rows)
    clips[k] = {'path': files[k], 'dups': dups, 'violation': any(org[d]['violation'] for d in dups),
                'org_times': o_times, 'human': h_rows, 'human_reliable': reliable,
                'layout': layout.get(k), 'group': find(k)}

os.makedirs(OUT, exist_ok=True)
json.dump(clips, open(os.path.join(OUT, 'gt.json'), 'w'), ensure_ascii=False, indent=1, default=str)

# ---- aggregate report only ------------------------------------------------------------
v = [c for c in clips.values() if c['violation']]
print(f'unique clips {len(clips)} | violation {len(v)} | clean {len(clips) - len(v)} | organizer events {sum(len(c["org_times"]) for c in v)}')
print(f'human violation rows {sum(len(h) for h in human.values())} on {len(human)} files | clean-sheet rows {clean_rows}')
print('human rows on organizer-clean clips:', sum(1 for c in clips.values() if c['human'] and not c['violation']))
print('organizer-violation clips without human row:', sum(1 for c in v if not c['human']))
agree = collections.Counter()
for c in v:
    ht = [t for h in c['human'] for t in h['times']]
    for t in c['org_times']:
        d = min((abs(t - x) for x in ht), default=None)
        agree['no human time' if d is None else '<=1s' if d <= 1 else '<=3s' if d <= 3 else '>3s'] += 1
print('organizer time vs nearest human time:', dict(agree))
print('human events per violation clip:', dict(collections.Counter(sum(max(1, len(h['times'])) for h in c['human']) for c in v)))
print('mechanism-reliable violation clips:', sum(c['human_reliable'] for c in v))
print('layout x verdict:', dict(collections.Counter((c['layout'], c['violation']) for c in clips.values())))
g = collections.Counter(c['group'] for c in clips.values())
print(f'groups {len(g)} | size histogram {dict(collections.Counter(g.values()))}')
print('norm1 x reliable:', dict(collections.Counter((h['norm1'], c['human_reliable']) for c in v for h in c['human'][:1])))
