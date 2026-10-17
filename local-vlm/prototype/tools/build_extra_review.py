"""Таблица для ручной проверки «лишних» событий: всё, что нашли прогоны, но чего нет в метках организаторов (±3 с).
Метки неполные: в ролике может быть несколько нарушений. Отмеченные «да» пойдут в brev_kit/data/gt_extra.csv,
и оценка перестанет считать их ложными срабатываниями.

python prototype/tools/build_extra_review.py --runs prototype/work/runs/proto_4b prototype/work/runs/proto_4b_v2 \
    --videos open/open/videos --out research/Доп_события_проверка.xlsx
"""
import argparse, csv, glob, json, os, subprocess, sys, tempfile

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image, ImageDraw

HERE = os.path.dirname(os.path.abspath(__file__))
KIT = os.path.join(os.path.dirname(os.path.dirname(HERE)), 'brev_kit')
sys.path.insert(0, os.path.join(KIT, 'scripts'))
from timeline_rules import TITLE, CODEX   # noqa: E402

# наблюдения пользователя, которых нет ни в метках, ни в предсказаниях
MANUAL = [('X-34', 7.0, 'no_id', 'наблюдение пользователя: клиент не показал удостоверение')]


def strip(vid, t, video, tmp):
    """Три кадра (t-1, t, t+1) в одну полоску 3×150 px с подписью времени."""
    tiles = []
    for dt in (-1, 0, 1):
        p = os.path.join(tmp, f'{vid}_{t + dt:.1f}.jpg')
        if subprocess.run(['ffmpeg', '-v', 'quiet', '-y', '-ss', f'{max(0, t + dt):.2f}', '-i', video, '-frames:v', '1', '-q:v', '3', p]).returncode \
                or not os.path.exists(p):
            continue
        im = Image.open(p).convert('RGB'); k = 150 / im.width; im = im.resize((150, round(im.height * k)))
        d = ImageDraw.Draw(im); d.rectangle((0, 0, 58, 16), fill=(0, 0, 0)); d.text((3, 2), f'{max(0, t + dt):.1f} s', fill=(255, 255, 255))
        tiles.append(im)
    if not tiles: return None
    out = Image.new('RGB', (150 * len(tiles) + 4 * (len(tiles) - 1), tiles[0].height), (255, 255, 255))
    for i, im in enumerate(tiles): out.paste(im, (i * 154, 0))
    path = os.path.join(tmp, f'strip_{vid}_{t:.1f}.png'); out.save(path); return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--runs', nargs='+', required=True); ap.add_argument('--videos', required=True); ap.add_argument('--out', required=True)
    a = ap.parse_args()
    gt = {}
    for r in csv.DictReader(open(os.path.join(KIT, 'data', 'gt_violations.csv'), encoding='utf-8'), delimiter=';'):
        gt.setdefault(r['video'], []).append((float(r['t']), r['type_ru']))
    pairs = {}
    for r in csv.DictReader(open(os.path.join(KIT, 'data', 'videos.csv'), encoding='utf-8'), delimiter=';'):
        pairs[r['video']] = r
    files = {os.path.splitext(os.path.basename(f))[0].replace('Х', 'X'): f for f in glob.glob(os.path.join(a.videos, '*'))}

    events = {}   # (vid, rule, t_bucket) -> dict
    def add(vid, t, rule, src, note=''):
        if any(abs(t - gt_t) <= 3 for gt_t, _ in gt.get(vid, [])) and not note: return     # совпало с меткой — не «лишнее»
        for (v, r, tt), e in events.items():
            if v == vid and r == rule and abs(tt - t) <= 3:
                e['src'].add(src); return
        events[(vid, rule, t)] = {'vid': vid, 't': t, 'rule': rule, 'src': {src}, 'note': note}

    for run in a.runs:
        name = os.path.basename(run.rstrip('/\\'))
        for p in glob.glob(os.path.join(run, 'violations', '*.json')):
            d = json.load(open(p, encoding='utf-8'))
            for v in d['violations']:
                add(d['video'], v['timecode_sec'], v['rule'], name + (' (не выдано: нужна проверка)' if v.get('needs_verify') else ''))
        mp = os.path.join(run, 'run_meta.json')
        if os.path.exists(mp):
            for vid, m in json.load(open(mp, encoding='utf-8')).get('videos', {}).items():
                for s in m.get('dropped', []):
                    rule, t = s.split('@'); add(vid, float(t), rule, name + ' (отсеяно подтверждением)')
    for vid, t, rule, note in MANUAL:
        add(vid, t, rule, 'вручную', note)

    wb = Workbook(); ws = wb.active; ws.title = 'Проверка'
    head = ['№', 'Видео', 'Секунда', 'Кадры (t−1, t, t+1)', 'Что нашла система', 'Норма', 'Метки организаторов в этом видео',
            'Источник', 'Реальное нарушение?', 'Комментарий']
    ws.append(['Дополнительные события: система нашла, а в метках их нет. Отметьте в колонке I: да / нет / не уверен. '
               'Подтверждённые «да» не будут считаться ложными срабатываниями.'])
    ws.append(head)
    for c in ws[2]: c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='DDEBF7'); c.alignment = Alignment(wrap_text=True, vertical='top')
    dv = DataValidation(type='list', formula1='"да,нет,не уверен"', allow_blank=True); ws.add_data_validation(dv)
    widths = [5, 8, 9, 64, 34, 14, 34, 26, 14, 30]
    for i, w in enumerate(widths): ws.column_dimensions[chr(65 + i)].width = w
    order = sorted(events.values(), key=lambda e: (int(e['vid'].split('-')[1]), e['t']))
    with tempfile.TemporaryDirectory() as tmp:
        for n, e in enumerate(order, 1):
            row = ws.max_row + 1
            labels = '; '.join(f'{t:g} с — {ty}' for t, ty in gt.get(e['vid'], [])) or 'чистое видео (нарушений нет)'
            ws.append([n, e['vid'], e['t'], '', TITLE.get(e['rule'], e['rule']) + (f" ({e['note']})" if e['note'] else ''),
                       ', '.join(CODEX.get(e['rule'], [])), labels, '; '.join(sorted(e['src'])), 'да' if e['src'] == {'вручную'} else '', ''])
            for c in ws[row]: c.alignment = Alignment(wrap_text=True, vertical='top')
            dv.add(ws.cell(row=row, column=9))
            if e['vid'] in files:
                p = strip(e['vid'], e['t'], files[e['vid']], tmp)
                if p:
                    img = XLImage(p); ws.add_image(img, f'D{row}'); ws.row_dimensions[row].height = img.height * 0.75 + 4
        wb.save(a.out)
    print('rows', len(order), '->', a.out)


if __name__ == '__main__':
    main()
