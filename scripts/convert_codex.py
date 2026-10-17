"""Convert the supplied workbook without authoring legal text or citations."""

import argparse
import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from backend.clip_id import parse_clip_id  # noqa: E402
from backend.f1.codex import (  # noqa: E402
    ClipInfo,
    Codex,
    Rule,
    duration_seconds,
    load_codex,
    rules_markdown,
)


def records(sheet) -> list[dict]:
    rows = iter(sheet.iter_rows(values_only=True))
    headers = next(rows)
    return [dict(zip(headers, row, strict=True)) for row in rows if any(v is not None for v in row)]


def convert(source: Path, destination: Path, rules_doc: Path) -> Codex:
    template_path = Path(__file__).resolve().parents[1] / "config/rule_templates.yaml"
    templates = yaml.safe_load(template_path.read_text(encoding="utf-8"))
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        for name in ("Cases", "Rules", "Case_Rule", "Clips"):
            if name not in workbook.sheetnames:
                raise ValueError(f"Missing worksheet: {name}")
        links: dict[str, list[str]] = {}
        for row in records(workbook["Case_Rule"]):
            links.setdefault(str(row["Case ID"]), []).append(str(row["Rule ID"]))
        rules = []
        for row in records(workbook["Rules"]):
            text = str(row["Min duration (proposal)"] or "")
            rules.append(
                Rule(
                    id=row["Rule ID"],
                    family=row["Family"],
                    requirement=row["Must be true"],
                    breach_condition=row["Breach condition (engine)"],
                    signals=row["Signals needed"],
                    min_duration_text=text,
                    min_duration_sec=duration_seconds(text),
                    timecode_anchor=row["Timecode anchor"],
                    source=row["Source (short)"],
                    source_url=row["Source link"],
                    strength=row["Strength"],
                    status=row["Status"],
                    title=templates.get(row["Rule ID"], {}).get("title", row["Rule ID"]),
                    description_template=templates.get(row["Rule ID"], {}).get("description", ""),
                    title_en=templates.get(row["Rule ID"], {}).get("title_en", ""),
                    description_template_en=templates.get(row["Rule ID"], {}).get(
                        "description_en", ""
                    ),
                )
            )
        clips = [
            ClipInfo(
                id=parse_clip_id(row["Clip"]),
                camera_layout=row["Camera layout"],
                rule_ids=links.get(row["Case ID (primary)"], []),
                case_id=row["Case ID (primary)"],
                duplicate_of=parse_clip_id(row["Duplicate of"]) if row["Duplicate of"] else None,
            )
            for row in records(workbook["Clips"])
        ]
        codex = Codex(source_workbook=source.name, rules=rules, clips=clips)
    finally:
        workbook.close()
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(
        yaml.safe_dump(codex.model_dump(), allow_unicode=True, sort_keys=False), encoding="utf-8"
    )
    codex = load_codex(destination)
    rules_doc.parent.mkdir(parents=True, exist_ok=True)
    rules_doc.write_text(rules_markdown(codex), encoding="utf-8")
    return codex


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=Path("data/Cash_Discipline_Cases_Rules.xlsx"))
    parser.add_argument("--out", type=Path, default=Path("config/f1_codex.yaml"))
    parser.add_argument("--rules-md", type=Path, default=Path("docs/F1_RULES.md"))
    args = parser.parse_args()
    codex = convert(args.input, args.out, args.rules_md)
    print(f"Converted {len(codex.rules)} rules and {len(codex.clips)} clip mappings")


if __name__ == "__main__":
    main()
